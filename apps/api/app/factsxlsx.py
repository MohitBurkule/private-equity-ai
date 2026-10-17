"""Excel export of verified facts. One row per fact; each value cell carries a comment listing its
sources as `filename p.N: quote`, so a reader can trace every number without opening the app.
Only facts passing the normal visibility gate are included."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.acl import fact_visible
from app.models import Document, Fact

HEADERS = ["Entity", "Metric", "Period", "Value", "Unit", "Basis", "Status", "Reviewed by"]
MAX_COMMENT = 800  # Excel comments get unreadable beyond this; full quotes stay in the app


def visible_facts(db: Session, tenant_id: str, entity: str | None, scope: set[str] | None) -> list[Fact]:
    q = db.query(Fact).filter_by(tenant_id=tenant_id)
    if entity:
        q = q.filter_by(entity=entity)
    # Only verified facts leave the app as a workbook; pending OCR and rejected rows stay in the queue.
    return [f for f in q.filter_by(status="verified").order_by(Fact.entity, Fact.metric, Fact.period, Fact.id).all()
            if fact_visible(f, scope)]


def fact_rows(db: Session, tenant_id: str, entity: str | None, scope: set[str] | None) -> list[list]:
    """HEADERS then one row per visible fact, values as plain cells (for the Google Sheets push)."""
    return [HEADERS] + [[f.entity, f.metric, f.period, float(f.value), f.unit, f.basis, f.status, f.reviewed_by or ""]
                        for f in visible_facts(db, tenant_id, entity, scope)]


def build_facts_xlsx(db: Session, tenant_id: str, entity: str | None, scope: set[str] | None) -> bytes:
    facts = visible_facts(db, tenant_id, entity, scope)
    names = {d.id: d.filename for d in db.query(Document).filter(Document.tenant_id == tenant_id)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Facts"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, f in enumerate(facts, start=2):
        ws.append([f.entity, f.metric, f.period, float(f.value), f.unit, f.basis, f.status, f.reviewed_by or ""])
        cited = [f"{names.get(s.document_id, s.document_id)} p.{s.page}: {s.quote}" for s in f.sources]
        if cited:
            text = "\n".join(cited)
            if len(text) > MAX_COMMENT:
                text = text[: MAX_COMMENT - 3] + "..."
            ws.cell(row=i, column=4).comment = Comment(text, "pe-ai")
    for col, width in zip("ABCDEFGH", [28, 28, 14, 16, 10, 8, 14, 16]):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
