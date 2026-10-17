"""Monthly KPI workbook for a portfolio company: each row becomes a verified, cited fact and then a KPI observation.

Layout: a sheet with a header row containing Metric, Period, Value and Unit (any order, any case). Period is
YYYY-MM. Metric names are matched against KPI_MAP only; anything else is reported as skipped, never guessed.
Each fact cites its sheet (page), the value's cell (anchor), and the row's text (quote). observed_at is the time
the pack arrives, so a monthly pack resets the KPI cadence watcher."""
from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.facts import CitationError, reject_unit_mix
from app.ingest import ingest_xlsx
from app.models import Fact, FactSource
from app.portfolio import PortfolioCompany, record_observation

KPI_MAP = {  # workbook label (lower case) -> KPI metric name
    "revenue": "revenue", "net revenue": "revenue",
    "ebitda": "ebitda", "adjusted ebitda": "ebitda",
    "gross margin": "gross_margin", "headcount": "headcount",
    "cash": "cash", "net debt": "net_debt",
}
_PERIOD = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
_HEADERS = ("metric", "period", "value", "unit")


@dataclass
class PackResult:
    observations: list[int]
    skipped: list[dict]
    unchanged: bool = False


def _header(ws) -> tuple[int, dict[str, int]] | None:
    for r_idx, row in enumerate(ws.iter_rows(max_row=30), 1):
        cols = {str(c.value).strip().lower(): c.column for c in row if isinstance(c.value, str)}
        if all(h in cols for h in _HEADERS):
            return r_idx, {h: cols[h] for h in _HEADERS}
    return None


def ingest_kpi_pack(db: Session, tenant_id: str, company_id: int, data: bytes, filename: str) -> PackResult:
    company = db.get(PortfolioCompany, company_id)
    if company is None or company.tenant_id != tenant_id:
        raise ValueError("company not found")
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "pack.xlsx"
        path.write_bytes(data)
        wb = load_workbook(path, data_only=True, read_only=True)
        doc = ingest_xlsx(db, tenant_id, path, filename)
        if any(f.document_id == doc.id for f in db.query(FactSource).filter_by(document_id=doc.id).limit(1)):
            return PackResult([], [], unchanged=True)
        out = PackResult([], [])
        for sheet_no, ws in enumerate(wb.worksheets, 1):
            found = _header(ws)
            if found is None:
                continue
            hdr_row, cols = found
            for r_idx, row in enumerate(ws.iter_rows(min_row=hdr_row + 1), hdr_row + 1):
                cells = {k: row[c - 1].value for k, c in cols.items() if c - 1 < len(row)}
                if all(v is None for v in cells.values()):
                    continue
                where = f"{ws.title}!{ws.cell(r_idx, cols['value']).coordinate}"
                label = str(cells["metric"] or "").strip().lower()
                metric = KPI_MAP.get(label)
                period = str(cells["period"] or "").strip()
                unit = str(cells["unit"] or "").strip()
                try:
                    value = Decimal(str(cells["value"]).replace(",", ""))
                except (InvalidOperation, ValueError):
                    value = None
                if metric is None or not _PERIOD.match(period) or value is None or not unit:
                    out.skipped.append({"cell": where, "reason": f"unmapped or unreadable row: {label!r} {period!r}"})
                    continue
                quote = f"{where}, row: {cells['metric']} | {period} | {cells['value']} {unit}"
                try:
                    reject_unit_mix(db, tenant_id, company.name, metric, period, "MGMT", unit)
                except CitationError as e:
                    out.skipped.append({"cell": where, "reason": str(e)})
                    continue
                fact = Fact(tenant_id=tenant_id, entity=company.name, metric=metric, period=period, value=value,
                            unit=unit, basis="MGMT", status="verified")
                fact.sources = [FactSource(document_id=doc.id, page=sheet_no, quote=quote, kind="xlsx",
                                           heading=ws.title, anchor=where.split("!", 1)[1])]
                db.add(fact)
                db.commit()
                out.observations.append(record_observation(db, tenant_id, company_id, fact.id).id)
        wb.close()
    return out
