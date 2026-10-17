"""LBO returns model. Pure code; every input must carry a citation.

Required inputs: entry EBITDA, entry and exit multiples, leverage (x EBITDA), interest rate, tax rate,
cash conversion, fees, hold period. Optional inputs, only when cited: EBITDA growth, revolver limit and
rate, management equity share, dividend recap year and target leverage.

Debt: a revolver funds cash shortfalls up to its limit and is repaid first; the term loan takes the rest
of free cash flow. A shortfall beyond the revolver limit is added to the term loan and flagged.
Dividend recap re-levers the term loan to the target multiple of that year's EBITDA and pays the gain to
the sponsor. Management equity takes a share of exit equity only.
Returns are the sponsor's: equity in at close; dividends and exit proceeds out. The sensitivity grid
re-runs the same model at other entry and exit multiples; those are scenarios, not cited inputs.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal

REQUIRED = ("ebitda", "entry_multiple", "exit_multiple", "leverage", "interest_rate",
            "tax_rate", "cash_conversion", "fees_pct", "years")
OPTIONAL = ("ebitda_growth", "revolver_limit", "revolver_rate", "mgmt_equity_pct", "recap_year", "recap_leverage")
D0 = Decimal(0)


@dataclass
class Cited:
    value: Decimal
    source: dict  # document_id, filename, page, quote


@dataclass
class Result:
    entry_ev: Decimal
    debt: Decimal
    equity: Decimal
    fees: Decimal
    schedule: list[dict] = field(default_factory=list)
    exit_ev: Decimal = D0
    exit_net_debt: Decimal = D0
    exit_equity: Decimal = D0  # before management equity
    dividends: Decimal = D0
    moic: Decimal = D0
    irr: Decimal | None = None


def missing(inputs: dict[str, Cited]) -> list[str]:
    return [k for k in REQUIRED if k not in inputs]


def _irr(cfs: list[Decimal]) -> Decimal | None:
    """Bisection on NPV of yearly cash flows, cfs[0] at t=0. None when no sign change."""
    if not (cfs[0] < 0 < sum(cfs[1:])):
        return None
    lo, hi = Decimal("-0.99"), Decimal(10)
    for _ in range(200):
        mid = (lo + hi) / 2
        npv = sum(cf / (1 + mid) ** t for t, cf in enumerate(cfs))
        if npv > 0:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


def run(inputs: dict[str, Cited]) -> Result:
    gap = missing(inputs)
    if gap:
        raise ValueError(f"missing inputs: {', '.join(gap)}")
    v = {k: c.value for k, c in inputs.items()}
    years = int(v["years"])
    if years < 1 or years > 10:
        raise ValueError("years must be 1 to 10")
    growth = v.get("ebitda_growth", D0) / 100
    rev_limit = v.get("revolver_limit", D0)
    rev_rate = v.get("revolver_rate", D0) / 100
    mgmt = v.get("mgmt_equity_pct", D0) / 100
    recap_year = int(v["recap_year"]) if "recap_year" in v else None
    if recap_year is not None and not 1 <= recap_year <= years:
        raise ValueError("recap_year must fall within the hold")
    if recap_year is not None and "recap_leverage" not in v:
        raise ValueError("recap_year needs recap_leverage")

    entry_ev = v["ebitda"] * v["entry_multiple"]
    term = v["ebitda"] * v["leverage"]  # leverage is a multiple of EBITDA
    fees = entry_ev * v["fees_pct"] / 100
    equity = entry_ev + fees - term
    if equity <= 0:
        raise ValueError("debt covers the whole price; equity must be positive")
    res = Result(entry_ev, term, equity, fees)
    revolver = D0
    cfs = [-equity] + [D0] * years
    ebitda = v["ebitda"]
    for y in range(1, years + 1):
        ebitda = ebitda * (1 + growth) if y > 1 else ebitda
        interest = term * v["interest_rate"] / 100 + revolver * rev_rate
        fcf = ebitda * v["cash_conversion"] / 100 - interest * (1 - v["tax_rate"] / 100)
        shortfall = False
        if fcf >= 0:
            repay = min(fcf, revolver)
            revolver -= repay
            term = max(term - (fcf - repay), D0)
        else:
            draw = min(-fcf, rev_limit - revolver)
            revolver += draw
            if -fcf > draw:
                term += -fcf - draw
                shortfall = True
        dividend = D0
        if recap_year == y:
            new_term = ebitda * v["recap_leverage"]
            dividend = max(new_term - term, D0)
            term = max(new_term, D0)
        cfs[y] += dividend
        res.dividends += dividend
        res.schedule.append({"year": y, "ebitda": str(ebitda.quantize(Decimal("0.01"))),
                             "interest": str(interest.quantize(Decimal("0.01"))),
                             "revolver_end": str(revolver.quantize(Decimal("0.01"))),
                             "term_end": str(term.quantize(Decimal("0.01"))),
                             "dividend": str(dividend.quantize(Decimal("0.01"))), "shortfall": shortfall})
    res.exit_ev = ebitda * v["exit_multiple"]
    res.exit_net_debt = term + revolver
    res.exit_equity = res.exit_ev - res.exit_net_debt
    cfs[years] += res.exit_equity * (1 - mgmt)
    res.moic = sum(cfs[1:]) / equity
    res.irr = _irr(cfs)
    return res


def sensitivity(inputs: dict[str, Cited], entry_multiples: list[Decimal],
                exit_multiples: list[Decimal]) -> list[list]:
    """Rows of [entry, exit, moic, irr]; each cell re-runs the model with that pair of multiples."""
    rows = []
    for en in entry_multiples:
        for ex in exit_multiples:
            trial = {**inputs,
                     "entry_multiple": Cited(en, inputs["entry_multiple"].source),
                     "exit_multiple": Cited(ex, inputs["exit_multiple"].source)}
            try:
                r = run(trial)
            except ValueError:
                rows.append([str(en), str(ex), None, None])
                continue
            rows.append([str(en), str(ex), str(r.moic.quantize(Decimal("0.01"))),
                         None if r.irr is None else str(r.irr.quantize(Decimal("0.0001")))])
    return rows


def build_xlsx(inputs: dict[str, Cited], res: Result, grid: list[list] | None = None) -> bytes:
    """Excel export: Inputs with a citation per input, the model, the schedule and the sensitivity grid."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append(["input", "value", "document", "page", "quote"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for k in REQUIRED + tuple(o for o in OPTIONAL if o in inputs):
        c = inputs[k]
        ws.append([k, str(c.value), c.source.get("filename"), c.source.get("page"), c.source.get("quote")])
    ws.column_dimensions["E"].width = 80
    out = wb.create_sheet("Model")
    out.append(["line", "value"])
    for line, val in [("entry EV", res.entry_ev), ("debt at close", res.debt), ("fees", res.fees),
                      ("sponsor equity", res.equity), ("exit EV", res.exit_ev),
                      ("net debt at exit", res.exit_net_debt), ("exit equity", res.exit_equity),
                      ("dividends to sponsor", res.dividends)]:
        out.append([line, str(val)])
    out.append(["MOIC", str(res.moic)])
    out.append(["IRR", str(res.irr) if res.irr is not None else "n/a"])
    out.append([])
    out.append(["year", "ebitda", "interest", "revolver_end", "term_end", "dividend", "shortfall"])
    for r in res.schedule:
        out.append([r["year"], r["ebitda"], r["interest"], r["revolver_end"], r["term_end"], r["dividend"],
                    r["shortfall"]])
    if grid:
        s = wb.create_sheet("Sensitivity")
        s.append(["entry multiple", "exit multiple", "MOIC", "IRR"])
        for row in grid:
            s.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
