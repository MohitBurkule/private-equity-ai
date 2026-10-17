"""Excel ingest: one 'page' per sheet, cell-addressed lines so citations point at cells.

Keeps formulas AND cached values, flags hidden sheets. Macros are never executed
(openpyxl reads only). Line format:  "Sheet1!B7: 1,234 [=SUM(B2:B6)]"
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class Sheet:
    doc_id: str
    number: int  # 1-based sheet index
    name: str
    hidden: bool
    text: str


def extract_sheets(path: Path) -> list[Sheet]:
    """Streams both workbooks in read-only mode: a 100 MB workbook is never held as a full object model.
    Values and formulas are read as two row streams in step, so a cell's formula is the one in the same
    position, without random access into the formula sheet."""
    doc_id = hashlib.sha256(path.read_bytes()).hexdigest()
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    try:
        out = []
        for i, (ws, fws) in enumerate(zip(values.worksheets, formulas.worksheets), 1):
            lines = []
            for vrow, frow in zip(ws.iter_rows(min_row=1, min_col=1), fws.iter_rows(min_row=1, min_col=1)):
                label = next((str(c.value) for c in vrow if isinstance(c.value, str) and c.value.strip()), None)
                for vc, fc in zip(vrow, frow):
                    if vc.value is None and fc.value is None:
                        continue
                    coord = vc.coordinate
                    f = fc.value
                    is_formula = isinstance(f, str) and f.startswith("=")
                    suffix = f" [{f}]" if is_formula else ""
                    if vc.value is None:  # never silently drop: workbook was saved without calculation
                        v = "(no cached value)"
                    else:
                        v = f"{vc.value:,}" if isinstance(vc.value, (int, float)) else str(vc.value)
                    row_note = f" | row: {label}" if label and label != str(vc.value) else ""
                    lines.append(f"{ws.title}!{coord}: {v}{suffix}{row_note}")
            hidden = ws.sheet_state != "visible"
            head = f"[hidden sheet: {ws.sheet_state}]\n" if hidden else ""
            out.append(Sheet(doc_id, i, ws.title, hidden, head + "\n".join(lines)))
        return out
    finally:
        values.close()
        formulas.close()
