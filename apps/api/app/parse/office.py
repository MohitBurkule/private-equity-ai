"""DOCX and PPTX text extraction into citeable 'pages'. Nothing is executed (macros ignored).

DOCX has no fixed pages, so body blocks (paragraphs and tables, in document order) are packed into
chunks of about CHUNK chars; each chunk is labelled with the nearest heading. Tables become
`row N: a | b | c` lines. PPTX is one page per slide (hidden slides flagged, speaker notes kept).
"""
from __future__ import annotations

import io
from dataclasses import dataclass

CHUNK = 3000


@dataclass(frozen=True)
class RawPage:
    text: str
    label: str | None = None


def _table_lines(table) -> list[str]:
    out = []
    for n, row in enumerate(table.rows, 1):
        cells = []
        for c in row.cells:
            t = " ".join(c.text.split())
            if not cells or t != cells[-1]:  # merged cells repeat; keep the first
                cells.append(t)
        if any(cells):
            out.append(f"row {n}: " + " | ".join(cells))
    return out


def extract_docx(data: bytes) -> list[RawPage]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(io.BytesIO(data))
    blocks: list[tuple[str, str | None]] = []  # (text, heading_if_heading)
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            t = p.text.strip()
            if t:
                heading = t if (p.style is not None and p.style.name.lower().startswith(("heading", "title"))) else None
                blocks.append((t, heading))
        elif tag == "tbl":
            for ln in _table_lines(Table(child, doc)):
                blocks.append((ln, None))
    pages: list[RawPage] = []
    buf: list[str] = []
    size = 0
    heading = None
    start_heading = None
    for text, h in blocks:
        if h:
            heading = h
        if not buf:
            start_heading = heading
        buf.append(text)
        size += len(text) + 1
        if size >= CHUNK:
            pages.append(RawPage("\n".join(buf), start_heading))
            buf, size = [], 0
    if buf:
        pages.append(RawPage("\n".join(buf), start_heading))
    return pages


def extract_pptx(data: bytes) -> list[RawPage]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        if slide._element.get("show") == "0":
            lines.append("[HIDDEN SLIDE]")
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else None
        for sh in slide.shapes:
            if sh.has_text_frame:
                lines += [ln.strip() for ln in sh.text_frame.text.splitlines() if ln.strip()]
            if getattr(sh, "has_table", False) and sh.has_table:
                lines += _table_lines(sh.table)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Notes: " + notes)
        pages.append(RawPage("\n".join(lines) or "[no text on this slide]", title or f"Slide {i}"))
    return pages
