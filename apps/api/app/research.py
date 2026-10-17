"""Target research note (Word): one page per target built only from stored records and their source links.

Sections: identity and where it came from, revenue with its filing link (or "not stated"), sector benchmark,
prior mentions across the firm's own records, and a Sources list. Nothing is estimated: a missing field is
written as "not stated in the record", so a reader can see the gap instead of a made-up value.
"""
from __future__ import annotations

import io

from sqlalchemy.orm import Session

from app import seen_before as seen_mod
from app import targets as tg_mod
from app.models import Target


def _none(v) -> str:
    return "not stated in the record" if v in (None, "") else str(v)


def build_docx(db: Session, tenant: str, t: Target, scope_deals: set[str] | None) -> bytes:
    from docx import Document as Docx

    d = Docx()
    d.add_heading(t.name, level=1)
    d.add_paragraph("DRAFT research note. Built from stored records only; each figure links to its source.")

    d.add_heading("Identity", level=2)
    for label, val in (("Source", t.source), ("Source reference", t.source_ref), ("Sector", t.sector),
                       ("SIC", t.sic), ("Country", t.country), ("Region", t.region), ("Pipeline status", t.status)):
        d.add_paragraph(f"{label}: {_none(val)}", style="List Bullet")

    d.add_heading("Revenue", level=2)
    if t.revenue is None:
        d.add_paragraph("Revenue: not stated in the record. No estimate is shown.")
    else:
        money = " ".join(x for x in (t.currency, str(t.revenue)) if x)
        d.add_paragraph(f"Revenue: {money} for {_none(t.revenue_period)}.")
        d.add_paragraph(f"Filing: {_none(t.revenue_url)}", style="List Bullet")

    d.add_heading("Sector benchmark", level=2)
    b = tg_mod.benchmark(db, tenant, t.sector.split(" (")[0] if t.sector else None, t)
    if b["peers"] == 0:
        d.add_paragraph("No peers with stated revenue in this sector. No benchmark shown.")
    else:
        d.add_paragraph(f"{b['peers']} peer(s) with stated revenue, not counting this target; "
                        f"{b['excluded_no_revenue']} excluded for no revenue.")
        d.add_paragraph(f"Median {b['median']}; interquartile {b['q1']} to {b['q3']}.", style="List Bullet")
        if "target" in b:
            d.add_paragraph(f"This target sits at percentile {b['target']['percentile']} among those peers.",
                            style="List Bullet")

    d.add_heading("Prior mentions in firm records", level=2)
    hits = seen_mod.lookup(db, tenant, t.name, scope_deals)["hits"]
    others = [h for h in hits if not (h["kind"] == "target" and h["id"] == t.id)]
    if not others:
        d.add_paragraph("No other matching record. This may be a first look; confirm by name variants.")
    for h in others:
        d.add_paragraph(f"{h['kind'].replace('_', ' ')}: {h['name']} ({h['rule']} name match)", style="List Bullet")

    d.add_heading("Sources", level=2)
    srcs = [("Target record", f"{t.source}:{t.source_ref}", t.source_url)]
    if t.revenue_url:
        srcs.append(("Revenue filing", t.revenue_period or "period not stated", t.revenue_url))
    for label, ref, url in srcs:
        d.add_paragraph(f"{label}: {ref} {url or ''}".strip(), style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
