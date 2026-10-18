"""QoE EBITDA bridge from an Excel sheet: each row becomes a candidate bridge line, cited to its cell.

Layout: a sheet with a header row containing Item, Kind, Amount and Period (any order, any case). Kind must be one of
reported, add_back, deduction, adjusted. Anything else is reported as skipped, never guessed, because a bridge line
with the wrong kind moves normalized EBITDA. Lines come back as candidates only: a reviewer accepts each one before it
counts (qoe.normalized_ebitda reads verified lines only).
"""
from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.facts import CitationError, reject_unit_mix
from app.ingest import ingest_xlsx
from app.models import Fact, FactSource

KINDS = ("reported", "add_back", "deduction", "adjusted")
_HEADERS = ("item", "kind", "amount", "period")


def _header(ws) -> tuple[int, dict[str, int]] | None:
    for r_idx, row in enumerate(ws.iter_rows(max_row=30), 1):
        cols = {str(c.value).strip().lower(): c.column for c in row if isinstance(c.value, str)}
        if all(h in cols for h in _HEADERS):
            return r_idx, {h: cols[h] for h in _HEADERS}
    return None


def bridge_lines(data: bytes) -> dict:
    """{"lines": [...], "skipped": [...]}. Each line: sheet, cell (the amount's address), item, kind, period, amount
    (Decimal), and a quote in the workbook form the citation check reads: "Sheet!C5: 20 | row: Interest expense"."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    lines, skipped = [], []
    for sheet_no, ws in enumerate(wb.worksheets, 1):
        head = _header(ws)
        if head is None:
            continue
        r_idx, col = head
        for row in ws.iter_rows(min_row=r_idx + 1):
            cells = {c.column: c for c in row}
            item = cells.get(col["item"])
            amount = cells.get(col["amount"])
            if item is None or item.value is None or amount is None or amount.value is None:
                continue
            where = f"{ws.title}!{amount.coordinate}"
            kind = str(cells[col["kind"]].value or "").strip().lower() if col["kind"] in cells else ""
            try:
                value = Decimal(str(amount.value))
            except (InvalidOperation, TypeError):
                skipped.append({"sheet": ws.title, "cell": amount.coordinate, "reason": "amount is not a number"})
                continue
            if kind not in KINDS:
                skipped.append({"sheet": ws.title, "cell": amount.coordinate, "reason": f"kind {kind!r} not one of {KINDS}"})
                continue
            period = str(cells[col["period"]].value or "").strip() if col["period"] in cells else ""
            if not period:
                skipped.append({"sheet": ws.title, "cell": amount.coordinate, "reason": "no period"})
                continue
            label = str(item.value).strip()
            lines.append({"sheet": ws.title, "cell": amount.coordinate, "item": label, "kind": kind,
                          "period": period, "amount": value, "page": sheet_no,
                          "quote": f"{where}: {amount.value} | row: {label}"})
    return {"lines": lines, "skipped": skipped}


def store_candidates(db: Session, tenant_id: str, entity: str, path: Path, filename: str) -> dict:
    """Stores a bridge workbook as a cited document and its bridge lines as candidate facts (status auto, basis
    MGMT). Candidates never count: qoe.normalized_ebitda reads verified lines only, so a reviewer accepts each line
    first. A line that would mix units with a verified figure is skipped and reported, never stored."""
    doc = ingest_xlsx(db, tenant_id, path, filename)
    parsed = bridge_lines(path.read_bytes())
    stored, skipped = [], list(parsed["skipped"])
    for line in parsed["lines"]:
        metric = f"ebitda_bridge {line['kind']}"
        try:
            reject_unit_mix(db, tenant_id, entity, metric, line["period"], "MGMT", "USD")
        except CitationError as e:
            skipped.append({"sheet": line["sheet"], "cell": line["cell"], "reason": str(e)})
            continue
        fact = Fact(tenant_id=tenant_id, entity=entity, metric=metric, period=line["period"], value=line["amount"],
                    unit="USD", basis="MGMT", status="auto")
        fact.sources = [FactSource(document_id=doc.id, page=line["page"], quote=line["quote"], kind="xlsx",
                                   heading=line["sheet"], anchor=line["cell"])]
        db.add(fact)
        stored.append(fact)
    db.commit()
    return {"document_id": doc.id, "fact_ids": [f.id for f in stored], "skipped": skipped}
