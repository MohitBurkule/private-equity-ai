"""Model hygiene checks on an Excel model's formulas: hardcoded numbers inside formula rows, references that
resolve to nothing (#REF!), and links to other workbooks.

Why a separate check: a LBO model that reads correctly can still be wrong in ways a reviewer only finds by hand:
a typed-over projection, a deleted row that left #REF! behind, or a formula that points at a file that is not in
the room. Each finding names its cell so a reviewer can go straight to it. A finding is a prompt, not a verdict:
a hardcoded input in a formula row is sometimes deliberate.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

MIN_FORMULAS_IN_ROW = 2  # a row with one formula is usually a single total, not a projection to check


def audit(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=False)
    out: list[dict] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None]
            formulas = [c for c in cells if isinstance(c.value, str) and c.value.startswith("=")]
            for c in cells:
                ref = f"{ws.title}!{c.coordinate}"
                if isinstance(c.value, str) and c.value.startswith("="):
                    if "#REF!" in c.value:
                        out.append({"cell": ref, "kind": "broken_reference", "detail": c.value})
                    if "[" in c.value and "]" in c.value:
                        out.append({"cell": ref, "kind": "external_link", "detail": c.value})
                elif isinstance(c.value, (int, float)) and len(formulas) >= MIN_FORMULAS_IN_ROW:
                    out.append({"cell": ref, "kind": "hardcode_in_formula_row", "detail": str(c.value)})
    return out
