"""IC memo section builder. Every number gets a footnote to document, page, quote.

Derived numbers are computed here in code from cited facts, never by a model.
Only status="verified" facts appear (OCR-sourced and unconfirmed facts are excluded).
"""
from __future__ import annotations

import io
import itertools
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal

from sqlalchemy.orm import Session

from app.acl import fact_visible
from app.citations import locator_label
from app.firm_knowledge import applied_rules
from app.models import Document, Fact


def _fmt(v: Decimal) -> str:
    return f"{v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):,}"


def _gloss(v: Decimal) -> str:
    """Exact figure first; large values get a short reading aid, e.g. '74,039,000,000.00 (about 74.04B)'."""
    a = abs(v)
    for size, tag in ((Decimal(10) ** 12, "T"), (Decimal(10) ** 9, "B"), (Decimal(10) ** 6, "M")):
        if a >= size:
            return f"{_fmt(v)} (about {(v / size).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}{tag})"
    return _fmt(v)


@dataclass
class Memo:
    entity: str
    lines: list[str] = field(default_factory=list)  # bullet text, includes [n] markers
    derived: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # footnote texts, index = marker - 1
    rules: list[dict] = field(default_factory=list)  # accepted firm rules this draft applies, each quoted
    series: dict[str, list[tuple[str, Decimal, int]]] = field(default_factory=dict)  # metric -> (period, value, note)
    risks: list[str] = field(default_factory=list)  # only what cited figures show, each with its note numbers

    @property
    def empty(self) -> bool:
        return not self.lines



# (label, numerator metric, denominator metric) for margins computed from cited facts
MARGINS = (("gross margin", "gross_profit", "revenue"), ("operating margin", "operating_income", "revenue"),
           ("net margin", "net_income", "revenue"))


def derive_lines(facts: dict[tuple[str, str], tuple[Decimal, str, int]]) -> list[str]:
    """Margins, free cash flow and year-on-year revenue growth, each citing the note numbers of its inputs.
    Pure: `facts` maps (metric, period) to (value, unit, note number). Nothing is derived across units."""
    out: list[str] = []
    periods = sorted({p for _, p in facts})
    for label, num, den in MARGINS:
        for period in periods:
            if (num, period) in facts and (den, period) in facts:
                nv, nu, nn = facts[(num, period)]
                dv, du, dn = facts[(den, period)]
                if dv and nu == du:
                    out.append(f"{label} ({period}): {_fmt(nv / dv * 100)}% = {num} [{nn}] / {den} [{dn}]")
    for period in periods:
        if ("operating_cash_flow", period) in facts and ("capex", period) in facts:
            ov, ou, on = facts[("operating_cash_flow", period)]
            cv, cu, cn = facts[("capex", period)]
            if ou == cu:
                out.append(f"free cash flow ({period}): {_fmt(ov - abs(cv))} {ou} = operating_cash_flow [{on}] "
                           f"- capex [{cn}]")
    for prev, cur in itertools.pairwise(periods):
        if ("revenue", prev) in facts and ("revenue", cur) in facts:
            pv, pu, pn = facts[("revenue", prev)]
            cv, cu, cn = facts[("revenue", cur)]
            if pv and pu == cu:
                out.append(f"revenue growth ({cur} vs {prev}): {_fmt((cv - pv) / pv * 100)}% = revenue [{cn}] "
                           f"vs revenue [{pn}]")
    return out


def risk_lines(facts: dict[tuple[str, str], tuple[Decimal, str, int]]) -> list[str]:
    """Risks the cited figures themselves show: losses, falling revenue, negative free cash flow."""
    out: list[str] = []
    for (metric, period), (v, _unit, n) in sorted(facts.items(), key=lambda kv: kv[0][1]):
        if metric in ("net_income", "operating_income") and v < 0:
            out.append(f"{metric.replace('_', ' ')} negative in {period} ({_fmt(v)}) [{n}]")
    periods = sorted({p for _, p in facts})
    for prev, cur in itertools.pairwise(periods):
        if ("revenue", prev) in facts and ("revenue", cur) in facts:
            pv, pu, pn = facts[("revenue", prev)]
            cv, cu, cn = facts[("revenue", cur)]
            if pu == cu and cv < pv:
                out.append(f"revenue fell in {cur} vs {prev} [{cn}, {pn}]")
    for period in periods:
        if ("operating_cash_flow", period) in facts and ("capex", period) in facts:
            ov, _, on = facts[("operating_cash_flow", period)]
            cv, _, cn = facts[("capex", period)]
            if ov - abs(cv) < 0:
                out.append(f"free cash flow negative in {period} [{on}, {cn}]")
    return out

def build_memo_data(db: Session, tenant_id: str, entity: str, scope: set[str] | None = None) -> Memo:
    memo = Memo(entity)
    facts = (db.query(Fact).filter_by(tenant_id=tenant_id, entity=entity, status="verified")
             .order_by(Fact.period, Fact.metric).all())
    by_key: dict[tuple[str, str], Fact] = {}
    note_no: dict[int, int] = {}
    facts = [f for f in facts if fact_visible(f, scope)]
    for f in facts:
        by_key[(f.metric, f.period)] = f
        n = len(memo.notes) + 1
        note_no[f.id] = n
        src = f.sources[0]
        doc = db.get(Document, src.document_id)
        where = locator_label(src.page, src.section, src.heading)
        memo.notes.append(f'[{n}] {doc.filename if doc else src.document_id[:8]}, {where}: '
                          f'"{src.quote.strip()}"')
        memo.lines.append(f"{f.metric} ({f.period}, {f.basis}): {_gloss(f.value)} {f.unit} [{n}]")
    memo.derived += derive_lines({k: (f.value, f.unit, note_no[f.id]) for k, f in by_key.items()})
    for (metric, period), f in sorted(by_key.items(), key=lambda kv: kv[0][1]):
        memo.series.setdefault(metric, []).append((period, f.value, note_no[f.id]))
    memo.risks = risk_lines({k: (f.value, f.unit, note_no[f.id]) for k, f in by_key.items()})
    memo.rules = applied_rules(db, tenant_id, scope)
    return memo


def _rule_line(r: dict) -> str:
    where = locator_label(r["source"]["page"])
    return f'{r["label"]}: {r["value"]} [{r["source"]["document_id"][:8]} {where}: "{r["source"]["quote"].strip()}"]'


def build_memo(db: Session, tenant_id: str, entity: str, scope: set[str] | None = None) -> str:
    m = build_memo_data(db, tenant_id, entity, scope)
    if m.empty and not m.rules:
        return f"# {entity}\n\nNo verified facts available. Nothing to report.\n"
    out = [f"# {entity}: financial summary (draft, analyst review required)", ""]
    out += [f"- {ln}" for ln in m.lines]
    if m.derived:
        out += ["", "Derived (computed in code):", *[f"- {d}" for d in m.derived]]
    if m.rules:
        out += ["", "Firm rules applied (accepted, quoted):", *[f"- {_rule_line(r)}" for r in m.rules]]
    return "\n".join([*out, "", "Sources:", *m.notes]) + "\n"


def build_memo_docx(db: Session, tenant_id: str, entity: str, scope: set[str] | None = None) -> bytes:
    from docx import Document as Docx
    m = build_memo_data(db, tenant_id, entity, scope)
    d = Docx()
    d.add_heading(f"{entity}: financial summary", level=1)
    d.add_paragraph("DRAFT. Analyst review required before use. Figures below are only those whose "
                    "quote and number were verified against the cited source page.")
    if m.empty:
        d.add_paragraph("No verified facts available. Nothing to report.")
    for ln in m.lines:
        d.add_paragraph(ln, style="List Bullet")
    if m.derived:
        d.add_heading("Derived (computed in code)", level=2)
        for ln in m.derived:
            d.add_paragraph(ln, style="List Bullet")
    if m.rules:
        d.add_heading("Firm rules applied", level=2)
        for r in m.rules:
            d.add_paragraph(_rule_line(r), style="List Bullet")
    if m.notes:
        d.add_heading("Sources", level=2)
        for n in m.notes:
            d.add_paragraph(n)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_memo_pptx(db: Session, tenant_id: str, entity: str, scope: set[str] | None = None) -> bytes:
    """IC deck: title, figures (8 per slide, each with its footnote marker), derived numbers, sources.
    Same verified-only data as the text and Word memo."""
    from pptx import Presentation
    from pptx.util import Pt
    m = build_memo_data(db, tenant_id, entity, scope)
    prs = Presentation()

    def bullets(title: str, items: list[str], size: int = 16):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.clear()
        for i, t in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = t
            para.font.size = Pt(size)

    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = f"{entity}: financial summary"
    t.placeholders[1].text = "DRAFT. Analyst review required. Only figures whose quote and number were verified against the cited page."
    if m.empty:
        bullets("No verified facts", ["Nothing to report."])
    for i in range(0, len(m.lines), 8):
        bullets("Figures" + (f" ({i // 8 + 1})" if len(m.lines) > 8 else ""), m.lines[i:i + 8])
    if m.rules:
        bullets("Firm rules applied", [_rule_line(r) for r in m.rules], size=12)
    if m.derived:
        bullets("Derived (computed in code)", m.derived)
    rev = m.series.get("revenue", [])
    if len(rev) >= 2:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = "Revenue trend (cited figures only)"
        cd = CategoryChartData()
        cd.categories = [p for p, _, _ in rev]
        cd.add_series("Revenue (USD)", [float(v) for _, v, _ in rev])
        sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(9), Inches(5), cd)
    bullets("Risks (from cited figures)",
            (m.risks or ["No loss, revenue decline or negative free cash flow in the cited years."])
            + ["Not in this draft: legal red flags and precedent deals (no source documents supplied for this run)."], 14)
    for i in range(0, len(m.notes), 6):
        bullets("Sources" + (f" ({i // 6 + 1})" if len(m.notes) > 6 else ""), m.notes[i:i + 6], 11)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
