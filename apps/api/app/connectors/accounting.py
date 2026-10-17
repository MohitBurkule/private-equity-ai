"""Portfolio-company accounting: monthly trial balances from a ledger's API (NetSuite, QuickBooks or a similar
system), reduced to revenue and net income per month and rendered into the 'Monthly' workbook template
(app/accounting.py). The sync stores a normal workbook, so the cited-fact reader works on it unchanged.

opts: base_url (https, required), company (required), periods (comma list of YYYY-MM, required).
Token is a bearer token, used per request, never stored. The API shape is SYNTHETIC until a real ledger's docs
are in hand:
  GET {base_url}/v1/companies/{company}/trial-balance?period=YYYY-MM
      -> {"period": "YYYY-MM", "accounts": [{"name", "type": "income" | "expense" | "other", "amount_usd"}]}
Amounts are positive as the ledger reports them. Net income is income minus expense; other accounts are ignored.
"""
from __future__ import annotations

import hashlib
import json
import re
from urllib.parse import quote

from openpyxl import Workbook

from app.connectors.base import Connector, ConnectorError, Item
from app.connectors.workbook import save_stable

_PERIOD = re.compile(r"\d{4}-\d{2}")


def monthly_row(tb: dict) -> tuple[str, float, float]:
    period = str(tb["period"])
    income = sum(float(a["amount_usd"]) for a in tb["accounts"] if a["type"] == "income")
    expense = sum(float(a["amount_usd"]) for a in tb["accounts"] if a["type"] == "expense")
    return period, round(income, 2), round(income - expense, 2)


def render_workbook(rows: list[tuple[str, float, float]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly"
    ws["A1"], ws["B1"], ws["C1"] = "Period", "Revenue (USD)", "Net income (USD)"
    for i, (period, revenue, net) in enumerate(rows, 2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = period, revenue, net
    return save_stable(wb)


class PortcoAccounting(Connector):
    name, label = "accounting", "Portfolio-company accounting"

    def items(self, limit: int = 50):
        base = self.opts.get("base_url", "").rstrip("/")
        company = self.opts.get("company", "")
        periods = [p.strip() for p in self.opts.get("periods", "").split(",") if p.strip()]
        if not base.startswith("https://") or not company or not periods:
            raise ConnectorError("opts.base_url (https), opts.company and opts.periods are required")
        if any(not _PERIOD.fullmatch(p) for p in periods):
            raise ConnectorError("opts.periods must be YYYY-MM months")
        headers = {"Authorization": f"Bearer {self.token}", "Accept": "application/json"}
        root = f"{base}/v1/companies/{quote(company, safe='')}"
        rows = []
        try:
            for p in periods[:limit]:
                rows.append(monthly_row(json.loads(self.http(f"{root}/trial-balance?period={p}", headers))))
        except (KeyError, ValueError, TypeError) as e:
            raise ConnectorError(f"accounting response not in the expected shape: {e}") from e
        data = render_workbook(rows)
        yield Item(hashlib.sha256(root.encode()).hexdigest()[:24], f"{company}-monthly.xlsx", data,
                   f"{base}/", hashlib.sha256(data).hexdigest()[:16])
