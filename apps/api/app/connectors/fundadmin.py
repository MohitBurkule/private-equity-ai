"""Fund administrator API: capital calls and NAV for one fund, rendered into the fund-admin workbook template
(app/fundadmin.py), so the sync stores a normal workbook and the cited-fact reader works on it unchanged.

opts: base_url (https, required), fund (required). Token is a bearer token, used per request, never stored.
The API shape assumed here is SYNTHETIC until a real administrator's documentation is in hand:
  GET {base_url}/v1/funds/{fund}/capital-calls -> {"items": [{"call_no", "date", "amount_usd"}]}
  GET {base_url}/v1/funds/{fund}/nav           -> {"as_of", "nav_usd"}
"""
from __future__ import annotations

import hashlib
import io
import json
from urllib.parse import quote

from openpyxl import Workbook

from app.connectors.base import Connector, ConnectorError, Item


def render_workbook(calls: list[dict], nav: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Capital Calls"
    ws["A1"], ws["B1"], ws["C1"] = "Call", "Date", "Amount (USD)"
    for i, c in enumerate(calls, 2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = c["call_no"], c["date"], c["amount_usd"]
    n = wb.create_sheet("NAV")
    n["A1"], n["B1"] = "As of", "NAV (USD)"
    n["A2"], n["B2"] = nav["as_of"], nav["nav_usd"]
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class FundAdminApi(Connector):
    name, label = "fundadmin", "Fund administrator API"

    def items(self, limit: int = 50):
        base = self.opts.get("base_url", "").rstrip("/")
        fund = self.opts.get("fund", "")
        if not base.startswith("https://") or not fund:
            raise ConnectorError("opts.base_url (https) and opts.fund are required")
        headers = {"Authorization": f"Bearer {self.token}", "Accept": "application/json"}
        root = f"{base}/v1/funds/{quote(fund, safe='')}"
        try:
            calls = json.loads(self.http(f"{root}/capital-calls", headers))["items"]
            nav = json.loads(self.http(f"{root}/nav", headers))
        except (KeyError, ValueError, TypeError) as e:
            raise ConnectorError(f"fund admin response not in the expected shape: {e}") from e
        data = render_workbook(calls[:limit], nav)
        digest = hashlib.sha256(data).hexdigest()
        yield Item(hashlib.sha256(root.encode()).hexdigest()[:24], f"{fund}-fund-admin.xlsx", data,
                   f"{base}/", digest[:16])
