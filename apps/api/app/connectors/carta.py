"""Carta cap table: the holdings of one company as of a date, rendered into the 'Cap Table' workbook template
(app/captable.py), so the sync stores a normal workbook and the cited-fact reader works on it unchanged.

opts: base_url (https, required), company (required), as_of (YYYY-MM-DD, required). Token is a bearer token, used
per request, never stored. The API shape is SYNTHETIC until Carta's documentation for the account is in hand:
  GET {base_url}/v1/companies/{company}/cap-table?as_of=YYYY-MM-DD
      -> {"holdings": [{"holder", "share_class", "shares", "ownership_pct"}]}
"""
from __future__ import annotations

import hashlib
import json
import re
from urllib.parse import quote

from openpyxl import Workbook

from app.connectors.base import Connector, ConnectorError, Item
from app.connectors.workbook import save_stable

_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")


def render_workbook(holdings: list[dict], as_of: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = "Holder", "Share class", "Shares", "Ownership (%)", "As of"
    for i, h in enumerate(holdings, 2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = h["holder"], h["share_class"], h["shares"]
        ws[f"D{i}"], ws[f"E{i}"] = h.get("ownership_pct"), as_of
    return save_stable(wb)


class Carta(Connector):
    name, label = "carta", "Carta cap table"

    def items(self, limit: int = 50):
        base = self.opts.get("base_url", "").rstrip("/")
        company = self.opts.get("company", "")
        as_of = self.opts.get("as_of", "")
        if not base.startswith("https://") or not company or not _DATE.fullmatch(as_of):
            raise ConnectorError("opts.base_url (https), opts.company and opts.as_of (YYYY-MM-DD) are required")
        headers = {"Authorization": f"Bearer {self.token}", "Accept": "application/json"}
        url = f"{base}/v1/companies/{quote(company, safe='')}/cap-table?as_of={as_of}"
        try:
            holdings = json.loads(self.http(url, headers))["holdings"]
            for h in holdings:
                h["holder"], h["share_class"], h["shares"] = str(h["holder"]), str(h["share_class"]), h["shares"]
        except (KeyError, ValueError, TypeError) as e:
            raise ConnectorError(f"carta response not in the expected shape: {e}") from e
        data = render_workbook(holdings[:limit], as_of)
        yield Item(hashlib.sha256(url.encode()).hexdigest()[:24], f"{company}-cap-table.xlsx", data,
                   f"{base}/", hashlib.sha256(data).hexdigest()[:16])
