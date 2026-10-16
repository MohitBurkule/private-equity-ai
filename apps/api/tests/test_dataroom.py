import io
import zipfile
from email.message import EmailMessage

from docx import Document as Docx
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def make_docx(extra="") -> bytes:
    d = Docx()
    d.add_heading("Revenue overview", level=1)
    d.add_paragraph("Total revenues were 1,234 in FY24. " + extra)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Metric", "FY24"
    t.cell(1, 0).text, t.cell(1, 1).text = "EBITDA", "321"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_pptx() -> bytes:
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Customer concentration"
    s.placeholders[1].text = "Top customer is 42% of revenue"
    s.notes_slide.notes_text_frame.text = "check with mgmt"
    h = p.slides.add_slide(p.slide_layouts[1])
    h.shapes.title.text = "Backup"
    h._element.set("show", "0")
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


def make_eml(att=None) -> bytes:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"], m["Date"] = "a@x.com", "b@y.com", "Q3 numbers", "Mon, 5 Oct 2026 10:00:00 +0000"
    m.set_content("Revenue was 5,000 in Q3.")
    if att:
        m.add_attachment(att, maintype="application", subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                         filename="memo.docx")
    return m.as_bytes()


def zipped(files: dict[str, bytes]) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        for k, v in files.items():
            z.writestr(k, v)
    return b.getvalue()


def post(h, name, data, **form):
    return c.post("/documents", files={"file": (name, data, "x")}, data=form, headers=h)


def page(h, doc, n=1):
    return c.get(f"/documents/{doc}/pages/{n}", headers=h).json()["text"]


def test_docx_text_headings_and_tables_are_citeable():
    h = new_tenant(c, "dr1")
    r = post(h, "cim.docx", make_docx()).json()
    t = page(h, r["id"])
    assert "Total revenues were 1,234" in t and "row 2: EBITDA | 321" in t
    fact = c.post("/facts", headers=h, json={"entity": "A", "metric": "ebitda", "period": "FY24", "value": "321",
                                             "unit": "USD", "document_id": r["id"], "page": 1,
                                             "quote": "row 2: EBITDA | 321"})
    assert fact.status_code == 201


def test_pptx_slides_notes_and_hidden_flag():
    h = new_tenant(c, "dr2")
    r = post(h, "deck.pptx", make_pptx()).json()
    assert r["n_pages"] == 2
    assert "42% of revenue" in page(h, r["id"], 1) and "Notes: check with mgmt" in page(h, r["id"], 1)
    assert "[HIDDEN SLIDE]" in page(h, r["id"], 2)


def test_eml_body_and_attachment_become_documents():
    h = new_tenant(c, "dr3")
    r = post(h, "thread.eml", make_eml(make_docx("from email"))).json()
    names = {d["filename"] for d in r["documents"]}
    assert names == {"thread.eml", "thread.eml/memo.docx"}
    assert "Subject: Q3 numbers" in page(h, r["id"]) and "Revenue was 5,000" in page(h, r["id"])


def test_zip_keeps_folder_tree_nested_zip_and_reports_skips():
    h = new_tenant(c, "dr4")
    wb = Workbook()
    wb.active["A1"] = "model"
    xb = io.BytesIO()
    wb.save(xb)
    inner = zipped({"deep/a.docx": make_docx("inner")})
    z = zipped({"Financials/2024/model.xlsx": xb.getvalue(), "Legal/cim.docx": make_docx(), "nested.zip": inner,
                "__MACOSX/junk": b"x", "notes.bin": b"\x00\x01\x02", "../evil.docx": make_docx("evil")})
    r = post(h, "dataroom.zip", z).json()
    names = {d["filename"] for d in r["documents"]}
    assert {"Financials/2024/model.xlsx", "Legal/cim.docx"} <= names
    assert any(n.endswith("deep/a.docx") for n in names)
    reasons = {s["name"]: s["reason"] for s in r["skipped"]}
    assert "path escapes archive" in reasons["../evil.docx"]
    assert "notes.bin" in reasons and "__MACOSX/junk" not in reasons


def test_corrupt_and_bomb_archives_rejected_without_crash(monkeypatch):
    h = new_tenant(c, "dr5")
    assert post(h, "broken.zip", b"PK\x03\x04garbage").status_code == 422
    import app.ingest as ing
    monkeypatch.setattr(ing, "MAX_ARCHIVE_BYTES", 10)
    r = post(h, "big.zip", zipped({"a.docx": make_docx()}))
    assert r.status_code == 422 and "zip bomb" in r.json()["detail"]


def test_deal_tag_applies_to_every_member_and_walls_hold():
    admin = {"X-Admin-Token": "test-admin"}
    t = c.post("/admin/tenants", json={"name": "dr6"}, headers=admin).json()
    owner = {"Authorization": f"Bearer {t['api_key']}"}
    k = c.post(f"/admin/tenants/{t['tenant_id']}/keys", json={"deals": ["alpha"]}, headers=admin).json()
    walled = {"Authorization": f"Bearer {k['api_key']}"}
    r = post(owner, "room.zip", zipped({"a.docx": make_docx("A"), "b.docx": make_docx("B")}), deal="bravo").json()
    assert len(r["documents"]) == 2
    assert c.get("/documents", headers=walled).json() == []
    assert all(d["deal"] == "bravo" for d in c.get("/documents", headers=owner).json())
    assert post(walled, "room2.zip", zipped({"c.docx": make_docx("C")}), deal="alpha").status_code == 200
    assert len(c.get("/documents", headers=walled).json()) == 1
