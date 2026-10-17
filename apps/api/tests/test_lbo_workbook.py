import io

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def _model():
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["B1"], ws["B2"] = 150, 9
    wb.defined_names["entry_ebitda"] = DefinedName("entry_ebitda", attr_text="Inputs!$B$1")
    wb.defined_names["entry_multiple"] = DefinedName("entry_multiple", attr_text="Inputs!$B$2")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_named_inputs_tie_to_verified_facts_or_are_flagged():
    h = new_tenant(c, "lbo-wb")
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = "EBITDA", 150
    b = io.BytesIO()
    wb.save(b)
    doc = c.post("/documents", files={"file": ("cim.xlsx", b.getvalue(), "x")}, headers=h).json()["id"]
    r = c.post("/facts", headers=h, json={"entity": "Co", "metric": "ebitda", "period": "FY24", "value": "150",
                                          "unit": "USD", "document_id": doc, "page": 1,
                                          "quote": "Sheet!B1: 150 | row: EBITDA"})
    assert r.status_code == 201, r.text
    out = c.post("/lbo/workbook/ties", headers=h,
                 files={"file": ("model.xlsx", _model(), "x")}).json()
    assert out["inputs"]["entry_ebitda"]["tie"] == "tied"
    assert out["inputs"]["entry_multiple"]["tie"] == "untied"
    assert out["untied"] == ["entry_multiple"]
    assert "tax_rate" in out["lbo_missing"]


def test_same_value_but_different_concept_is_not_a_tie():
    h = new_tenant(c, "lbo-wb2")
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = "Revenue", 9
    b = io.BytesIO()
    wb.save(b)
    doc = c.post("/documents", files={"file": ("rev.xlsx", b.getvalue(), "x")}, headers=h).json()["id"]
    c.post("/facts", headers=h, json={"entity": "Co", "metric": "revenue", "period": "FY24", "value": "9",
                                      "unit": "USD", "document_id": doc, "page": 1,
                                      "quote": "Sheet!B1: 9 | row: Revenue"})
    out = c.post("/lbo/workbook/ties", headers=h, files={"file": ("model.xlsx", _model(), "x")}).json()
    assert out["inputs"]["entry_multiple"]["tie"] == "untied"
