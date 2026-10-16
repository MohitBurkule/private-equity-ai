"""In-process end-to-end run (no Docker, no network): ingest -> cited fact -> memo -> exports -> report -> audit."""
import io

from docx import Document as Docx
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def _tiny_pdf() -> bytes:
    """One page whose text layer says: Net income 150."""
    objs = [b"<</Type/Catalog/Pages 2 0 R>>", b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 400 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            None, b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"]
    stream = b"BT /F1 12 Tf 20 100 Td (Net income 150) Tj ET"
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        body = o or b"<</Length %d>>\nstream\n" % len(stream) + stream + b"\nendstream"
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    x = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    return out + b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)


def test_pipeline_end_to_end():
    h = new_tenant(c, "e2e")
    up = c.post("/documents", files={"file": ("a.pdf", _tiny_pdf(), "application/pdf")}, headers=h)
    assert up.status_code == 200, up.text
    doc = up.json()["id"]
    hit = c.get("/search", params={"q": "net income"}, headers=h).json()[0]
    fact = {"entity": "E2E", "metric": "net_income", "period": "FY", "value": "150", "unit": "USD",
            "document_id": doc, "page": hit["page"], "quote": "Net income 150"}
    assert c.post("/facts", json=fact, headers=h).status_code == 201
    assert c.post("/facts", json={**fact, "value": "151"}, headers=h).status_code == 422
    memo = c.get("/memo", params={"entity": "E2E"}, headers=h).json()["markdown"]
    assert "[1]" in memo and "150" in memo
    d = c.get("/memo.docx", params={"entity": "E2E"}, headers=h)
    assert d.status_code == 200 and "150" in "\n".join(p.text for p in Docx(io.BytesIO(d.content)).paragraphs)
    p = c.get("/memo.pptx", params={"entity": "E2E"}, headers=h)
    assert p.status_code == 200 and len(Presentation(io.BytesIO(p.content)).slides) >= 1
    assert c.get("/portfolio/report", headers=h).status_code == 200
    assert c.get("/audit/verify", headers=h).json()["ok"] is True
