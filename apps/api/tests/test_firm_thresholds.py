import io
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app, llm_dependency
from tests.helpers import new_tenant

c = TestClient(app)


class Fake:
    def __init__(self, values):
        self.values = values

    def complete(self, system, prompt):
        return json.dumps({"values": self.values})


def upload(h, rows):
    wb = Workbook()
    for i, (k, v) in enumerate(rows, 1):
        wb.active.cell(i, 1, k)
        if v is not None:
            wb.active.cell(i, 2, v)
    b = io.BytesIO()
    wb.save(b)
    return c.post("/documents", files={"file": ("p.xlsx", b.getvalue(), "x")}, headers=h).json()["id"]


def fact(h, doc, metric, period, value, quote, unit="x"):
    r = c.post("/facts", headers=h, json={"entity": "Co", "metric": metric, "period": period, "value": str(value),
                                          "unit": unit, "document_id": doc, "page": 1, "quote": quote})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_leverage_limit_from_accepted_firm_rule_alerts_on_breach_only():
    h = new_tenant(c, "ft1")
    doc = upload(h, [("Maximum net leverage covenant 4.5x", None), ("Net leverage Q1 3.9x", None), ("Net leverage Q2 5.2x", None)])
    # the firm rule: extracted, quote-checked, then accepted by a human
    app.dependency_overrides[llm_dependency] = lambda: Fake([{
        "key": "leverage_limit", "value": "4.5x maximum net leverage", "page": 1,
        "quote": "Maximum net leverage covenant 4.5x"}])
    try:
        r = c.post("/firm-knowledge/extract", headers=h, json={"document_id": doc}).json()
    finally:
        app.dependency_overrides.pop(llm_dependency)
    assert len(r["saved"]) == 1, r["rejected"]
    c.post(f"/firm-knowledge/{r['saved'][0]['id']}/review", headers=h, json={"status": "accepted"})
    co = c.post("/portfolio/companies", json={"name": "Co"}, headers=h).json()["id"]
    for per, v in (("Q1", 3.9), ("Q2", 5.2)):
        row = {"Q1": 2, "Q2": 3}[per]  # the workbook row this figure sits on
        o = fact(h, doc, "net_leverage_x", per, v, f"Sheet!A{row}: Net leverage {per} {v}x")
        assert c.post(f"/portfolio/companies/{co}/observations", params={"fact_id": o}, headers=h).status_code == 201
    made = c.post("/alerts/scan", headers=h).json()["created"]
    lev = [a for a in made if a["kind"] == "leverage"]
    assert len(lev) == 1 and "Q2" in lev[0]["title"] and "4.5" in lev[0]["title"]
    assert "Maximum net leverage covenant 4.5x" in lev[0]["detail"]  # the firm rule's own quote is shown
    assert c.post("/alerts/scan", headers=h).json()["created"] == []  # not repeated


def test_no_accepted_rule_means_no_leverage_alerts():
    h = new_tenant(c, "ft2")
    doc = upload(h, [("Net leverage Q1 9.9x", None)])
    co = c.post("/portfolio/companies", json={"name": "Co"}, headers=h).json()["id"]
    o = fact(h, doc, "net_leverage_x", "Q1", 9.9, "Sheet!A1: Net leverage Q1 9.9x")
    c.post(f"/portfolio/companies/{co}/observations", params={"fact_id": o}, headers=h)
    assert [a for a in c.post("/alerts/scan", headers=h).json()["created"] if a["kind"] == "leverage"] == []
