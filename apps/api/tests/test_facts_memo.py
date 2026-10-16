from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db import Base, make_engine
from app.facts import CitationError, add_verified_fact
from app.memo import build_memo
from app.models import Document, PageRow

PAGE = "Statement (in millions)\nTotal revenues 1,000\nNet income 150\n"


def setup_db():
    engine = make_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    s = Session(engine)
    d = Document(id="d" * 64, tenant_id="t", sha256="s", filename="acme.pdf", n_pages=1)
    d.pages = [PageRow(number=1, text=PAGE)]
    s.add(d)
    s.commit()
    return s, d.id


def test_fact_requires_matching_quote_and_number():
    s, doc = setup_db()
    kw = {"tenant_id": "t", "entity": "Acme", "period": "FY24", "unit": "USD", "document_id": doc, "page": 1}
    add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="Total revenues 1,000", **kw)
    with pytest.raises(CitationError):
        add_verified_fact(s, metric="revenue", value=Decimal(999), quote="Total revenues 1,000", **kw)
    with pytest.raises(CitationError):
        add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="Total sales 1,000", **kw)
    with pytest.raises(CitationError):
        add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="x", **{**kw, "page": 2})


def test_memo_has_footnotes_and_code_derived_margin():
    s, doc = setup_db()
    kw = {"tenant_id": "t", "entity": "Acme", "period": "FY24", "unit": "USD", "document_id": doc, "page": 1}
    add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="Total revenues 1,000", **kw)
    add_verified_fact(s, metric="net_income", value=Decimal(150), quote="Net income 150", **kw)
    md = build_memo(s, "t", "Acme")
    assert "[1]" in md and "acme.pdf, p.1" in md
    assert "net margin (FY24): 15.00%" in md


def test_memo_empty_reports_nothing():
    s, _ = setup_db()
    assert "Nothing to report" in build_memo(s, "t", "Nobody")


def test_api_roundtrip():
    from app.main import app
    from tests.helpers import new_tenant
    from tests.test_ingest import blank_pdf
    c = TestClient(app)
    h = new_tenant(c, "api-t")
    doc = c.post("/documents", files={"file": ("a.pdf", blank_pdf(1), "application/pdf")}, headers=h).json()
    r = c.post("/facts", json={"entity": "X", "metric": "revenue", "period": "FY", "value": "1",
                               "unit": "USD", "document_id": doc["id"], "page": 1, "quote": "nope"}, headers=h)
    assert r.status_code == 422
    assert c.get("/memo", params={"entity": "X"}, headers=h).status_code == 200
    assert c.get("/search", params={"q": "revenue"}, headers=h).status_code == 200


def test_docx_export_contains_figures_footnotes_and_draft_notice():
    import io

    from docx import Document as Docx

    from app.memo import build_memo_docx
    s, doc = setup_db()
    kw = {"tenant_id": "t", "entity": "Acme", "period": "FY24", "unit": "USD", "document_id": doc, "page": 1}
    add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="Total revenues 1,000", **kw)
    text = "\n".join(p.text for p in Docx(io.BytesIO(build_memo_docx(s, "t", "Acme"))).paragraphs)
    assert "DRAFT" in text and "revenue (FY24, GAAP): 1,000.00 USD [1]" in text
    assert '[1] acme.pdf, p.1: "Total revenues 1,000"' in text
    empty = "\n".join(p.text for p in Docx(io.BytesIO(build_memo_docx(s, "t", "Nobody"))).paragraphs)
    assert "Nothing to report" in empty


def test_docx_endpoint_requires_auth_and_returns_word_file():
    from app.main import app
    from tests.helpers import new_tenant
    c = TestClient(app)
    assert c.get("/memo.docx", params={"entity": "X"}).status_code == 401
    r = c.get("/memo.docx", params={"entity": "X"}, headers=new_tenant(c, "docx-t"))
    assert r.status_code == 200 and r.content[:2] == b"PK" and "wordprocessingml" in r.headers["content-type"]


def test_pptx_export_has_figures_sources_and_respects_empty():
    import io

    from pptx import Presentation

    from app.memo import build_memo_pptx
    s, doc = setup_db()
    kw = {"tenant_id": "t", "entity": "Acme", "period": "FY24", "unit": "USD", "document_id": doc, "page": 1}
    add_verified_fact(s, metric="revenue", value=Decimal(1000), quote="Total revenues 1,000", **kw)

    def texts(data):
        return "\n".join(sh.text_frame.text for sl in Presentation(io.BytesIO(data)).slides
                         for sh in sl.shapes if sh.has_text_frame)
    t = texts(build_memo_pptx(s, "t", "Acme"))
    assert "DRAFT" in t and "revenue (FY24, GAAP): 1,000.00 USD [1]" in t and '[1] acme.pdf, p.1: "Total revenues 1,000"' in t
    assert "Nothing to report" in texts(build_memo_pptx(s, "t", "Nobody"))


def test_pptx_endpoint_auth_and_type():
    from app.main import app
    from tests.helpers import new_tenant
    c = TestClient(app)
    assert c.get("/memo.pptx", params={"entity": "X"}).status_code == 401
    r = c.get("/memo.pptx", params={"entity": "X"}, headers=new_tenant(c, "pptx-t"))
    assert r.status_code == 200 and r.content[:2] == b"PK" and "presentationml" in r.headers["content-type"]


def test_gloss_adds_reading_aid_for_large_values():
    from decimal import Decimal
    from app.memo import _gloss
    assert _gloss(Decimal("74039000000")) == "74,039,000,000.00 (about 74.04B)"
    assert _gloss(Decimal("1000")) == "1,000.00"
