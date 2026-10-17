import io

from docx import Document as Docx
from fastapi.testclient import TestClient

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def test_research_note_states_gaps_and_sources():
    h = new_tenant(c, "rs1")
    c.post("/targets", headers=h, json=[
        {"name": "Acme Holdings, Inc.", "source_ref": "a", "sector": "Software (SIC 62012)", "revenue": "150",
         "currency": "USD", "revenue_period": "FY2023", "revenue_url": "http://sec/x"},
        {"name": "Peer One", "source_ref": "p1", "sector": "Software", "revenue": "50", "revenue_url": "http://x"},
        {"name": "Bare Co", "source_ref": "b", "sector": "Software"}])
    tid = next(t["id"] for t in c.get("/targets?q=Acme", headers=h).json()["items"])
    r = c.get(f"/targets/{tid}/research.docx", headers=h)
    assert r.status_code == 200
    text = "\n".join(p.text for p in Docx(io.BytesIO(r.content)).paragraphs)
    assert "Revenue: USD 150.00 for FY2023." in text and "http://sec/x" in text
    assert "1 peer(s) with stated revenue, not counting this target" in text
    assert "percentile" in text
    bare = next(t["id"] for t in c.get("/targets?q=Bare", headers=h).json()["items"])
    text = "\n".join(p.text for p in Docx(io.BytesIO(c.get(f"/targets/{bare}/research.docx", headers=h).content)).paragraphs)
    assert "Revenue: not stated in the record. No estimate is shown." in text
    other = new_tenant(c, "rs2")
    assert c.get(f"/targets/{tid}/research.docx", headers=other).status_code == 404


def test_research_deck_has_same_sections_and_source_links():
    from pptx import Presentation

    h = new_tenant(c, "rs3")
    c.post("/targets", headers=h, json=[
        {"name": "Deck Holdings", "source_ref": "d", "sector": "Software", "revenue": "90",
         "currency": "USD", "revenue_period": "FY2024", "revenue_url": "http://sec/deck"}])
    tid = next(t["id"] for t in c.get("/targets?q=Deck", headers=h).json()["items"])
    r = c.get(f"/targets/{tid}/research.pptx", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    prs = Presentation(io.BytesIO(r.content))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles[0] == "Research: Deck Holdings"
    assert titles[1:] == ["Identity", "Revenue", "Sector benchmark", "Prior mentions in firm records", "Sources"]
    body = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.placeholders if sh.has_text_frame)
    assert "Revenue: USD 90.00 for FY2024." in body and "http://sec/deck" in body
    other = new_tenant(c, "rs4")
    assert c.get(f"/targets/{tid}/research.pptx", headers=other).status_code == 404
