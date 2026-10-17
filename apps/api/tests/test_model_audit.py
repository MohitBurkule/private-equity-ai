import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def _xlsx(fill):
    wb = Workbook()
    ws = wb.active
    ws.title = "LBO"
    fill(ws)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_flags_hardcode_in_formula_row_broken_ref_and_external_link():
    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "EBITDA", "=B2*2", 150, "=C2+1"
        ws["B2"], ws["C2"] = 10, 5
        ws["A3"], ws["B3"] = "Broken", "=#REF!+1"
        ws["A4"], ws["B4"] = "Linked", "='[other.xlsx]Sheet1'!A1"
    h = new_tenant(c, "model1")
    out = c.post("/models/audit", headers=h, files={"file": ("m.xlsx", _xlsx(fill), "x")}).json()["findings"]
    kinds = {(f["cell"], f["kind"]) for f in out}
    assert ("LBO!C1", "hardcode_in_formula_row") in kinds
    assert ("LBO!B3", "broken_reference") in kinds
    assert ("LBO!B4", "external_link") in kinds
    assert not any(f["cell"] == "LBO!B2" for f in out)
