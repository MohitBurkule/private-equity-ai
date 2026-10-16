import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.citations import quote_in_page
from tests.helpers import new_tenant


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws["A1"], ws["B1"] = "Revenue", 1000
    ws["A2"], ws["B2"] = "Costs", 400
    ws["A3"], ws["B3"] = "EBITDA", "=B1-B2"
    hidden = wb.create_sheet("Scratch")
    hidden["A1"] = "old numbers"
    hidden.sheet_state = "hidden"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xlsx_upload_cell_addresses_formulas_and_hidden():
    from app.main import app
    c = TestClient(app)
    h = new_tenant(c, "xlsx-t")
    r = c.post("/documents", files={"file": ("m.xlsx", make_xlsx(), "application/octet-stream")}, headers=h)
    assert r.status_code == 200 and r.json()["n_pages"] == 2
    doc = r.json()["id"]
    p1 = c.get(f"/documents/{doc}/pages/1", headers=h).json()["text"]
    assert quote_in_page("P&L!B1: 1,000", p1)
    assert "[=B1-B2]" in p1
    assert "[hidden sheet: hidden]" in c.get(f"/documents/{doc}/pages/2", headers=h).json()["text"]


def test_unsupported_type_rejected():
    from app.main import app
    c = TestClient(app)
    r = c.post("/documents", files={"file": ("x.txt", b"hello", "text/plain")},
               headers=new_tenant(c, "txt-t"))
    assert r.status_code == 422


def test_formula_without_cached_value_is_flagged_not_dropped():
    from app.main import app
    c = TestClient(app)
    h = new_tenant(c, "nocache-t")
    doc = c.post("/documents", files={"file": ("m.xlsx", make_xlsx(), "x")}, headers=h).json()["id"]
    assert "P&L!B3: (no cached value) [=B1-B2]" in c.get(f"/documents/{doc}/pages/1", headers=h).json()["text"]
