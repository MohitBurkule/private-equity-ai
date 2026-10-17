import io
from decimal import Decimal

from docx import Document as Docx
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app import html_doc
from app.citations import locator_label
from app.db import Base, make_engine
from app.facts import add_html_fact, add_verified_fact
from app.factsxlsx import build_facts_xlsx
from app.memo import build_memo, build_memo_docx
from app.models import Document, HtmlUnit, PageRow

HTML = ('<html><body><p>2.1. Term Loan A Facility.</p>'
        '<p>7.12. Financial Covenants.</p>'
        '<p>Borrower shall not permit the Consolidated Total Net Leverage Ratio to be greater than 4.00:1.00.</p>'
        '</body></html>')


def _seed(s: Session) -> None:
    units = html_doc.parse_units(HTML.encode())
    s.add(Document(id="h" * 64, tenant_id="t", sha256="h", filename="tla.htm", n_pages=0, kind="html"))
    s.flush()
    s.add_all(HtmlUnit(tenant_id="t", document_id="h" * 64, idx=u.idx, section=u.section, heading=u.heading,
                       anchor=u.anchor, offset=u.offset, text=u.text) for u in units)
    s.add(Document(id="p" * 64, tenant_id="t", sha256="p", filename="acme.pdf", n_pages=4, kind="pdf"))
    s.flush()
    s.add(PageRow(document_id="p" * 64, number=4, text="Total revenues 1,000"))
    s.flush()
    cov = next(u for u in units if "Net Leverage" in u.text)
    add_html_fact(s, tenant_id="t", entity="XPO", metric="covenant total_net_leverage maximum", period="per agreement",
                  value=Decimal("4.00"), unit="x", document_id="h" * 64, sources=[(cov.idx, "greater than 4.00:1.00")])
    add_verified_fact(s, tenant_id="t", entity="XPO", metric="revenue", period="FY24", value=Decimal(1000),
                      unit="USD", document_id="p" * 64, page=4, quote="Total revenues 1,000")


def _session():
    engine = make_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    s = Session(engine)
    _seed(s)
    return s


def test_locator_label_never_shows_page_zero():
    assert locator_label(4) == "p.4"
    assert locator_label(0, "7.12", "Financial Covenants") == "§7.12 Financial Covenants"
    assert locator_label(0) == "HTML filing"
    assert "p.0" not in locator_label(0, None, None)


def test_no_output_prints_page_zero():
    s = _session()
    md = build_memo(s, "t", "XPO")
    docx_text = "\n".join(p.text for p in Docx(io.BytesIO(build_memo_docx(s, "t", "XPO"))).paragraphs)
    ws = load_workbook(io.BytesIO(build_facts_xlsx(s, "t", "XPO", None))).active
    xlsx_text = "\n".join(str(c.value) for r in ws.iter_rows() for c in r) + "".join(
        c.comment.text for r in ws.iter_rows() for c in r if c.comment)
    for text in (md, docx_text, xlsx_text):
        assert "p.0" not in text
    assert "§7.12 Financial Covenants" in md and "acme.pdf, p.4" in md
    assert "§7.12 Financial Covenants" in docx_text
    assert "§7.12 Financial Covenants" in xlsx_text
