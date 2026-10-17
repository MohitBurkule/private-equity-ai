from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.lbo import Cited, missing, run, sensitivity
from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)
SRC = {"document_id": "d1", "filename": "cim.pdf", "page": 4, "quote": "Adjusted EBITDA of $100.0 million"}
BASE = {"ebitda": "100", "entry_multiple": "10", "exit_multiple": "10", "leverage": "5", "interest_rate": "8",
        "tax_rate": "25", "cash_conversion": "60", "fees_pct": "2", "years": "5"}


def cited(d):
    return {k: Cited(Decimal(v), SRC) for k, v in d.items()}


def test_run_math_and_returns():
    r = run(cited(BASE))
    assert r.entry_ev == 1000 and r.debt == 500 and r.fees == 20 and r.equity == 520
    assert r.exit_ev == 1000 and r.exit_equity > 0 and r.moic > 1
    assert r.irr is not None and 0 < r.irr < 1
    assert len(r.schedule) == 5 and Decimal(r.schedule[-1]["term_end"]) < 500


def test_refuses_missing_and_bad_inputs():
    d = dict(BASE)
    del d["tax_rate"]
    assert missing(cited(d)) == ["tax_rate"]
    try:
        run(cited(d))
        raise AssertionError("expected ValueError")
    except ValueError as e:
        assert "tax_rate" in str(e)
    try:
        run(cited({**BASE, "leverage": "11"}))
        raise AssertionError("expected ValueError")
    except ValueError as e:
        assert "equity must be positive" in str(e)


def test_endpoint_requires_citation_and_exports_xlsx():
    h = new_tenant(c, "lbo1")
    inputs = {k: {"value": v, "source": SRC} for k, v in BASE.items()}
    ok = c.post("/lbo/returns", headers=h, json={"inputs": inputs})
    assert ok.status_code == 200 and ok.json()["equity"] == "520"
    bad = {**inputs, "ebitda": {"value": "100"}}
    assert c.post("/lbo/returns", headers=h, json={"inputs": bad}).status_code == 422
    short = {k: v for k, v in inputs.items() if k != "years"}
    assert c.post("/lbo/returns", headers=h, json={"inputs": short}).status_code == 422
    x = c.post("/lbo/returns?xlsx=true", headers=h, json={"inputs": inputs})
    assert x.status_code == 200
    import io
    wb = load_workbook(io.BytesIO(x.content))
    rows = list(wb["Inputs"].iter_rows(values_only=True))
    assert rows[1][0] == "ebitda" and rows[1][2] == "cim.pdf" and rows[1][3] == 4 and "Adjusted EBITDA" in rows[1][4]


def test_revolver_recap_and_mgmt_equity_change_returns():
    base = run(cited(BASE))
    growth = run(cited({**BASE, "ebitda_growth": "5"}))
    assert growth.exit_ev > base.exit_ev and growth.moic > base.moic
    mgmt = run(cited({**BASE, "mgmt_equity_pct": "10"}))
    assert mgmt.moic < base.moic and mgmt.irr is not None
    recap = run(cited({**BASE, "recap_year": "3", "recap_leverage": "6"}))
    assert recap.dividends > 0 and recap.schedule[2]["term_end"] == "600.00"
    # cash-negative year: revolver draws, is repaid first once cash turns positive
    neg = run(cited({**BASE, "cash_conversion": "5", "revolver_limit": "100", "revolver_rate": "10"}))
    assert Decimal(neg.schedule[0]["revolver_end"]) > 0 and neg.schedule[0]["shortfall"] is False
    tight = run(cited({**BASE, "cash_conversion": "5", "revolver_limit": "1"}))
    assert tight.schedule[0]["shortfall"] is True


def test_recap_needs_target_and_optional_bounds():
    try:
        run(cited({**BASE, "recap_year": "2"}))
        raise AssertionError("expected ValueError")
    except ValueError as e:
        assert "recap_leverage" in str(e)


def test_sensitivity_grid_rows_and_xlsx_sheet():
    grid = sensitivity(
        cited(BASE), [Decimal(9), Decimal(10)], [Decimal(10), Decimal(12)])
    assert len(grid) == 4 and grid[0][:2] == ["9", "10"]
    assert Decimal(grid[3][2]) > Decimal(grid[0][2])  # higher exit multiple, higher MOIC
    h = new_tenant(c, "lbo2")
    inputs = {k: {"value": v, "source": SRC} for k, v in BASE.items()}
    body = {"inputs": inputs, "entry_multiples": ["9", "10"], "exit_multiples": ["10", "12"]}
    js = c.post("/lbo/returns", headers=h, json=body).json()
    assert len(js["sensitivity"]) == 4
    x = c.post("/lbo/returns?xlsx=true", headers=h, json=body)
    import io
    assert load_workbook(io.BytesIO(x.content)).sheetnames == ["Inputs", "Model", "Sensitivity"]
    bad = c.post("/lbo/returns", headers=h, json={"inputs": {**inputs, "junk": {"value": "1", "source": SRC}}})
    assert bad.status_code == 422


def test_assumption_needs_who_when_why_and_is_listed_not_sourced():
    h = new_tenant(c, "lbo-assume")
    inputs = {k: {"value": v, "source": SRC} for k, v in BASE.items()}
    exit_ass = {"value": "11", "assumption": {"set_by": "deal team", "set_on": "2026-10-01",
                                              "reason": "in line with peer exits"}}
    body = {"inputs": {**inputs, "exit_multiple": exit_ass}}
    js = c.post("/lbo/returns", headers=h, json=body)
    assert js.status_code == 200
    out = js.json()
    assert out["assumptions"] == [{"input": "exit_multiple", "value": "11", "set_by": "deal team",
                                   "set_on": "2026-10-01", "reason": "in line with peer exits"}]
    # An assumption with no owner, date or reason is refused, the same as an uncited input.
    for missing_field in ("set_by", "set_on", "reason"):
        a = {k: v for k, v in exit_ass["assumption"].items() if k != missing_field}
        bad = {**inputs, "exit_multiple": {"value": "11", "assumption": a}}
        assert c.post("/lbo/returns", headers=h, json={"inputs": bad}).status_code == 422
    # A sourced run lists no assumptions.
    assert c.post("/lbo/returns", headers=h, json={"inputs": inputs}).json()["assumptions"] == []
