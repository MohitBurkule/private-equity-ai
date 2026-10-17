"""DOCX citations point at paragraph numbers: spans are stored per block and a verified quote maps through them."""
import io
import json

from docx import Document as Docx

from app.citations import paragraphs_for
from app.parse.office import CHUNK, extract_docx


def make_docx(paras: list[str], table: list[list[str]] | None = None) -> bytes:
    d = Docx()
    for t in paras:
        d.add_paragraph(t)
    if table:
        tbl = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                tbl.cell(r, c).text = val
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_spans_slice_back_to_their_block_text():
    data = make_docx(["Intro line.", "Revenue was 1,234 in FY24.", "Closing."], table=[["Item", "Value"], ["EBITDA", "99"]])
    (page,) = extract_docx(data)
    assert [n for _, _, n in page.paras] == [1, 2, 3, 4, 5]
    for start, end, n in page.paras:
        assert page.text[start:end].strip()
    s, e, n = page.paras[1]
    assert page.text[s:e] == "Revenue was 1,234 in FY24." and n == 2
    assert page.text[page.paras[3][0]:page.paras[3][1]] == "row 1: Item | Value"


def test_spans_restart_offsets_per_chunk_and_keep_document_numbers():
    long = "x" * (CHUNK // 2)
    data = make_docx([long, long, long, "Final figure 42."])
    pages = extract_docx(data)
    assert len(pages) > 1
    last = pages[-1]
    s, e, n = last.paras[-1]
    assert n == 4 and last.text[s:e] == "Final figure 42."


def test_quote_maps_to_its_paragraph_and_ranges_for_multi_paragraph_quotes():
    data = make_docx(["Intro line.", "Revenue was 1,234 in FY24.", "Closing."])
    (page,) = extract_docx(data)
    spans = json.dumps(list(page.paras))
    assert paragraphs_for("revenue was  1,234", page.text, spans) == [2]
    assert paragraphs_for("Intro line.\nRevenue was 1,234", page.text, spans) == [1, 2]


def test_no_spans_or_no_match_gives_no_paragraph():
    assert paragraphs_for("anything", "anything here", None) is None
    assert paragraphs_for("absent", "anything here", json.dumps([[0, 5, 1]])) is None
