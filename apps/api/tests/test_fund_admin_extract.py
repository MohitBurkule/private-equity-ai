"""Fund-admin workbook → cited facts with no typed-in /facts. SYNTHETIC workbook in the template layout; the
extractor reads the cell lines the xlsx ingest stored, so every fact's quote is text the citation check finds."""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import fundadmin
from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def _workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Capital Calls"
    ws["A1"], ws["B1"], ws["C1"] = "Call", "Date", "Amount (USD)"
    ws["A2"], ws["B2"], ws["C2"] = 4, "2026-09-30", 12_500_000
    ws["A3"], ws["B3"], ws["C3"] = 5, "2026-12-31", 9_000_000
    nav = wb.create_sheet("NAV")
    nav["A1"], nav["B1"] = "As of", "NAV (USD)"
    nav["A2"], nav["B2"] = "2026-06-30", 214_000_000
    other = wb.create_sheet("Notes")
    other["A1"], other["B1"] = "Comment", "Amount (USD)"
    other["A2"], other["B2"] = "ignore me", 1
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_extractor_reads_the_template_and_skips_other_sheets():
    text = ("Capital Calls!A1: Call | row: Call\nCapital Calls!B1: Date\nCapital Calls!C1: Amount (USD)\n"
            "Capital Calls!B2: 2026-09-30\nCapital Calls!C2: 12,500,000 | row: 2026-09-30")
    pages = [(1, text)]
    (found,) = fundadmin.extract(pages)
    assert (found.metric, found.period, found.value, found.page) == ("capital_call_usd", "2026-09-30", "12500000", 1)
    assert found.quote == "Capital Calls!C2: 12,500,000"


def test_workbook_becomes_cited_facts_and_citations_resolve():
    h = new_tenant(c, "fa-extract")
    doc = c.post("/documents", files={"file": ("fund7-admin-pack.xlsx", _workbook(), "x")}, headers=h).json()["id"]
    r = c.post(f"/documents/{doc}/fund-admin/facts", params={"entity": "Fund VII (SYNTHETIC)"}, headers=h)
    assert r.status_code == 201, r.text
    out = r.json()
    assert len(out["created"]) == 3 and out["skipped"] == []
    kpi_ids = []
    co = c.post("/portfolio/companies", json={"name": "Fund VII (SYNTHETIC)"}, headers=h).json()["id"]
    for fid in out["created"]:
        assert c.post(f"/portfolio/companies/{co}/observations?fact_id={fid}", headers=h).status_code == 201
        kpi_ids.append(fid)
    kpis = c.get(f"/portfolio/companies/{co}/kpis", headers=h).json()
    by_metric = sorted((k["metric"], k["period"], float(k["value"])) for k in kpis)
    assert by_metric == [("capital_call_usd", "2026-09-30", 12_500_000.0),
                         ("capital_call_usd", "2026-12-31", 9_000_000.0),
                         ("nav_usd", "2026-06-30", 214_000_000.0)]
    assert set(kpi_ids) == {k["fact_id"] for k in kpis}


def test_extraction_is_repeatable():
    h = new_tenant(c, "fa-repeat")
    doc = c.post("/documents", files={"file": ("pack.xlsx", _workbook(), "x")}, headers=h).json()["id"]
    first = c.post(f"/documents/{doc}/fund-admin/facts", params={"entity": "Fund VII"}, headers=h).json()
    again = c.post(f"/documents/{doc}/fund-admin/facts", params={"entity": "Fund VII"}, headers=h).json()
    assert sorted(first["created"]) == sorted(again["created"])
