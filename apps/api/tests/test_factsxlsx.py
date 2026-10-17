import io

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db import Base, make_engine
from app.factsxlsx import build_facts_xlsx
from app.models import Document, Fact, FactSource


def _seed(s: Session):
    s.add(Document(id="d1", tenant_id="t", sha256="a", filename="cim.pdf", n_pages=3, kind="pdf", deal="atlas"))
    s.add(Document(id="d2", tenant_id="t", sha256="b", filename="other.pdf", n_pages=3, kind="pdf", deal="secret"))
    s.flush()
    f1 = Fact(tenant_id="t", entity="Acme", metric="revenue", period="FY2024", value=1200, unit="USD_m",
              status="verified")
    f1.sources.append(FactSource(document_id="d1", page=4, quote="Revenue was $1,200m"))
    f2 = Fact(tenant_id="t", entity="Acme", metric="ebitda", period="FY2024", value=300, unit="USD_m")
    f2.sources.append(FactSource(document_id="d2", page=9, quote="EBITDA 300"))
    s.add_all([f1, f2])
    s.flush()


def test_facts_xlsx_values_comments_and_scope():
    engine = make_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        _seed(s)
        blob = build_facts_xlsx(s, "t", "Acme", scope={"d1"})
    ws = load_workbook(io.BytesIO(blob)).active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row)]
    assert rows[0][:4] == ["Entity", "Metric", "Period", "Value"]
    assert len(rows) == 2  # ebitda is cited only to d2, which is outside scope
    assert rows[1][:4] == ["Acme", "revenue", "FY2024", 1200.0]
    assert "cim.pdf p.4: Revenue was $1,200m" in ws["D2"].comment.text


def test_facts_xlsx_unrestricted_includes_all_and_empty_is_valid():
    engine = make_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        _seed(s)
        full = load_workbook(io.BytesIO(build_facts_xlsx(s, "t", None, None))).active
        empty = load_workbook(io.BytesIO(build_facts_xlsx(s, "t", "Nobody", None))).active
    assert full.max_row == 3
    assert empty.max_row == 1


def test_facts_xlsx_endpoint_returns_workbook():
    from fastapi.testclient import TestClient

    from app.main import app
    from tests.helpers import new_tenant
    h = new_tenant(TestClient(app), "xlsx")
    r = TestClient(app).get("/facts.xlsx", params={"entity": "Nobody"}, headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert load_workbook(io.BytesIO(r.content)).active["A1"].value == "Entity"
