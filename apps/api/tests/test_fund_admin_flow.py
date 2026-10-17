"""Fund-admin drop → portfolio KPIs, every number cited to its cell. The workbook is SYNTHETIC: a capital-call
register and a NAV line, as a fund administrator might send them. Each KPI keeps the fact it came from, and the
fact keeps the document and cell."""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from tests.helpers import new_tenant

c = TestClient(app)


def _workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Capital Calls"
    ws["A1"], ws["B1"], ws["C1"] = "Call", "Date", "Amount (USD)"
    ws["A2"], ws["B2"], ws["C2"] = 4, "2026-09-30", 12_500_000
    ws2 = wb.create_sheet("NAV")
    ws2["A1"], ws2["B1"] = "As of", "NAV (USD)"
    ws2["A2"], ws2["B2"] = "2026-06-30", 214_000_000
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_capital_call_and_nav_land_as_cited_portfolio_kpis():
    h = new_tenant(c, "fund-admin")
    doc = c.post("/documents", files={"file": ("fund7-admin-pack.xlsx", _workbook(), "x")}, headers=h).json()["id"]
    call = c.post("/facts", headers=h, json={"entity": "Fund VII (SYNTHETIC)", "metric": "capital_call_usd",
                                             "period": "2026-09-30", "value": "12500000", "unit": "USD",
                                             "document_id": doc, "page": 1,
                                             "quote": "Capital Calls!C2: 12,500,000"})
    nav = c.post("/facts", headers=h, json={"entity": "Fund VII (SYNTHETIC)", "metric": "nav_usd",
                                            "period": "2026-06-30", "value": "214000000", "unit": "USD",
                                            "document_id": doc, "page": 2, "quote": "NAV!B2: 214,000,000"})
    assert call.status_code == 201 and nav.status_code == 201, (call.text, nav.text)

    co = c.post("/portfolio/companies", json={"name": "Fund VII (SYNTHETIC)"}, headers=h).json()["id"]
    for fid in (call.json()["id"], nav.json()["id"]):
        assert c.post(f"/portfolio/companies/{co}/observations?fact_id={fid}", headers=h).status_code == 201

    kpis = {k["metric"]: k for k in c.get(f"/portfolio/companies/{co}/kpis", headers=h).json()}
    assert float(kpis["capital_call_usd"]["value"]) == 12_500_000 and float(kpis["nav_usd"]["value"]) == 214_000_000
    assert kpis["nav_usd"]["fact_id"] == nav.json()["id"]


def test_a_number_not_in_its_cell_is_refused():
    h = new_tenant(c, "fund-admin-bad")
    doc = c.post("/documents", files={"file": ("pack.xlsx", _workbook(), "x")}, headers=h).json()["id"]
    bad = c.post("/facts", headers=h, json={"entity": "Fund VII (SYNTHETIC)", "metric": "nav_usd", "period": "2026-06-30",
                                            "value": "300000000", "unit": "USD", "document_id": doc, "page": 2,
                                            "quote": "NAV!B2: 214,000,000"})
    assert bad.status_code == 422
