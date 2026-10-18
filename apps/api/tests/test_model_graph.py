"""Model graph: assumptions are typed numbers a formula reads; outputs are formulas nothing reads."""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.model_graph import graph
from tests.helpers import new_tenant

c = TestClient(app)


def _xlsx(fill):
    wb = Workbook()
    ws = wb.active
    ws.title = "LBO"
    fill(ws)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _model(ws):
    ws["A1"], ws["B1"] = "Entry multiple", 8
    ws["A2"], ws["B2"] = "EBITDA", 100
    ws["A3"], ws["B3"] = "Unused note", 42
    ws["C1"] = "=B1*B2"          # entry EV, read by the return line
    ws["C2"] = "=C1*0.5"          # read by the return line
    ws["D1"] = "=C2+SUM(B2:B2)"   # output: nothing reads it
    ws["E1"] = "=Other!A1*2"     # reads another sheet, which is not in this workbook; still an output


def test_assumptions_and_outputs_from_the_formula_graph():
    g = graph(_xlsx(_model))
    assumed = {a["cell"]: a["value"] for a in g["assumptions"]}
    assert assumed == {"LBO!B1": 8, "LBO!B2": 100}
    assert "LBO!B3" not in assumed
    outs = {o["cell"] for o in g["outputs"]}
    assert outs == {"LBO!D1", "LBO!E1"}


def test_endpoint_returns_graph_and_stores_nothing():
    h = new_tenant(c, "mg1")
    r = c.post("/models/graph", headers=h, files={"file": ("m.xlsx", _xlsx(_model), "x")})
    assert r.status_code == 200
    body = r.json()
    assert body["formula_cells"] == 4
    assert {a["cell"] for a in body["assumptions"]} == {"LBO!B1", "LBO!B2"}


def test_sheet_qualified_reference_counts_as_read():
    def fill(ws):
        ws["A1"] = 5
        ws["B1"] = "=LBO!A1*2"
        ws["C1"] = "=B1"
    g = graph(_xlsx(fill))
    assert [a["cell"] for a in g["assumptions"]] == ["LBO!A1"]
    assert [o["cell"] for o in g["outputs"]] == ["LBO!C1"]
